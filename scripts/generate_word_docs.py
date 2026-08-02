import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

output_dir = os.path.join(os.getcwd(), 'docs', 'rag_sample_word_docs')
os.makedirs(output_dir, exist_ok=True)

def create_docx(filename, title, sections):
    doc = docx.Document()
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    
    doc.add_paragraph() # Spacing
    
    for section_heading, paragraphs in sections:
        # Heading 2
        p_head = doc.add_paragraph()
        run_head = p_head.add_run(section_heading)
        run_head.font.name = 'Arial'
        run_head.font.size = Pt(14)
        run_head.font.bold = True
        run_head.font.color.rgb = RGBColor(79, 70, 229) # Indigo 600
        
        for p_text in paragraphs:
            p = doc.add_paragraph()
            run_p = p.add_run(p_text)
            run_p.font.name = 'Arial'
            run_p.font.size = Pt(11)
            run_p.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
            
        doc.add_paragraph() # Spacing
        
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    print(f"[OK] Created Word Document: {filepath}")

# 1. SLA Document
create_docx(
    "01_Customer_Support_SLA_and_Escalation.docx",
    "Enterprise Customer Support SLA & Escalation Policy",
    [
        (
            "1. Executive Overview & Availability Guarantees",
            [
                "Antigravity Enterprise provides 24/7 technical support, 99.95% availability commitments, and structured escalation workflows for all enterprise knowledge base services.",
                "System uptime is monitored continuously across all global endpoints. If monthly availability falls below 99.95%, qualified Enterprise accounts receive pro-rated service credits."
            ]
        ),
        (
            "2. SLA Priority Classification & Response Targets",
            [
                "Priority 1 (P1 - Critical Outage): Complete core system downtime affecting all users. Initial Response Target: Under 15 Minutes. Target Resolution Time: Under 2 Hours. Support Coverage: 24/7/365 Dedicated Phone & Slack Connect.",
                "Priority 2 (P2 - Major Degraded Service): Partial service degradation or latency exceeding 2000ms. Initial Response Target: Under 1 Hour. Target Resolution Time: Under 6 Hours. Support Coverage: 24/7 Email & Slack.",
                "Priority 3 (P3 - Moderate Issue): Non-critical bug or single user permission issue. Initial Response Target: Under 4 Business Hours. Target Resolution Time: Under 24 Hours. Support Coverage: Mon-Fri 8 AM - 8 PM EST."
            ]
        ),
        (
            "3. Multi-Tiered Incident Escalation Ladder",
            [
                "Level 1 Triage: Tier 1 Support Specialist collects error logs, status codes, and reproduction steps.",
                "Level 2 Technical Escalation: Tier 2 Engineer engaged after 30 minutes for P1 or 2 hours for P2 outages to debug database connection pools and router states.",
                "Level 3 Engineering: Tier 3 AI Solutions Architect applies code hotfixes, schema migrations, or provider overrides.",
                "Level 4 Executive Escalation: VP of Customer Operations (executive-escalations@antigravity.ai) receives direct alert for unresolved P1 outages past 90 minutes."
            ]
        )
    ]
)

# 2. Billing & Refund Document
create_docx(
    "02_Billing_Refund_and_Subscription_Policy.docx",
    "Subscription Billing, Payment, Refund & Data Retention Policy",
    [
        (
            "1. Subscription Tiers & Pricing Specs",
            [
                "Developer / Starter Plan ($29/month): Up to 50 uploaded documents (25MB max size) and 10,000 RAG queries per month. Includes standard email support.",
                "Pro Team Plan ($149/month): Up to 500 uploaded documents (100MB max size) and 100,000 RAG queries per month. Includes priority email and Slack support.",
                "Enterprise Custom Plan (Custom Quote): Unlimited uploaded documents and queries, 99.95% SLA, dedicated model router, and 24/7 phone support."
            ]
        ),
        (
            "2. 30-Day Money-Back Guarantee",
            [
                "All new subscription plans qualify for a 100% full money-back refund within 30 calendar days of initial purchase.",
                "Refund requests must be submitted via the Billing Portal or by emailing billing@antigravity.ai. Approved refunds are credited back to the original payment method within 3 to 5 business days."
            ]
        ),
        (
            "3. Cancellation & Automated GDPR / SOC 2 Data Purging",
            [
                "Subscribers may cancel their subscription anytime via Account Settings -> Subscription -> Cancel Plan. Access remains active through the current billing period.",
                "Grace Period (Days 0-30): Documents and vector embeddings remain safely archived for instant account reactivation.",
                "Permanent Purge (Day 30): All uploaded documents, parsed text chunks, vector embeddings, and telemetry logs are permanently and unrecoverably erased from all primary servers and database backups in compliance with GDPR Article 17 and SOC 2 Type II regulations."
            ]
        )
    ]
)

# 3. Security & Privacy Document
create_docx(
    "03_Account_Security_and_Privacy_Policy.docx",
    "Account Security, Access Control & Privacy Policy",
    [
        (
            "1. Role-Based Access Control (RBAC)",
            [
                "Customer Role: Authenticated or public users restricted exclusively to sending support queries to the chatbot interface. Cannot view vector stores, telemetry, or admin settings.",
                "Admin Role: Designated platform administrators with exclusive rights to manage knowledge base documents, configure router strategies, toggle OKF/RAG modes, and inspect analytics."
            ]
        ),
        (
            "2. Authentication & Password Protection",
            [
                "Multi-Factor Authentication (MFA): Mandatory for all admin accounts using TOTP authenticator apps or hardware keys.",
                "Password Security: Passwords must be at least 10 characters long with upper/lowercase letters, numbers, and symbols.",
                "Account Lockout: Accounts auto-lock for 15 minutes after 5 consecutive failed login attempts to prevent brute-force attacks."
            ]
        ),
        (
            "3. Data Encryption & Tenant Isolation",
            [
                "Encryption in Transit: All communications use TLS 1.3 encryption with ECDHE key exchanges.",
                "Encryption at Rest: Uploaded documents and vector embeddings are stored using AES-256 GCM encryption.",
                "Multi-Tenant Isolation: Database tenant boundaries are strictly enforced via PostgreSQL Row Level Security (RLS) policies."
            ]
        )
    ]
)

# 4. Returns & Warranty Document
create_docx(
    "04_Product_Return_Warranty_and_RMA_Guide.docx",
    "Product Return, Hardware Warranty & RMA Guidelines",
    [
        (
            "1. Return & Exchange Policy Window",
            [
                "30-Day Return Window: Customers may return unopened, undamaged, or defective hardware items within 30 days of delivery for a full refund or direct unit replacement.",
                "Items must be returned in original packaging with included power cables, mounting brackets, and documentation."
            ]
        ),
        (
            "2. 1-Year Limited Hardware Warranty",
            [
                "All hardware units carry a 1-Year Limited Warranty covering manufacturing defects, power supply failures, and component breakdowns under normal operating conditions.",
                "Compliant with the US FTC Magnuson-Moss Warranty Act and European Union Consumer Protection Directives."
            ]
        ),
        (
            "3. Return Merchandise Authorization (RMA) Steps",
            [
                "Step 1: Contact support or email rma@antigravity.ai with the product serial number and description of the issue.",
                "Step 2: Our logistics team issues an official RMA tracking code and prepaid shipping label within 4 business hours.",
                "Step 3: Securely package the unit and drop off at an authorized carrier shipment center.",
                "Step 4: Once scanned by the carrier, a replacement unit is dispatched via 2-day expedited air delivery."
            ]
        )
    ]
)

# 5. Customer Troubleshooting Manual
create_docx(
    "05_Customer_Troubleshooting_and_FAQ_Manual.docx",
    "Customer Technical Support & Troubleshooting Manual",
    [
        (
            "1. Account & Password Troubleshooting",
            [
                "Password Reset: Click 'Forgot Password' on the login screen, enter your email, and follow the link sent to your inbox within 15 minutes.",
                "Account Unlocking: If locked out due to failed attempts, wait 15 minutes or click 'Unlock via Email' to receive an instant unlock link."
            ]
        ),
        (
            "2. Latency & Connection Troubleshooting",
            [
                "Latency Spikes (>2000ms): Transient latency triggers automatic multi-model router failovers. If slow responses persist past 5 minutes, clear your browser cache or test on an alternate network.",
                "Slack Connect Support: Pro and Enterprise customers can connect their Slack team via Account Settings -> Support -> Slack Connect."
            ]
        ),
        (
            "3. General Inquiries & Feature Requests",
            [
                "Support Hours: 24/7/365 for P1 critical issues; Mon-Fri 8 AM - 8 PM EST for general technical support.",
                "Feature Suggestions: Submit feature requests via the Support Portal under Submit Ticket -> Request Feature."
            ]
        )
    ]
)

print("\n[SUCCESS] All 5 Enterprise Customer Support Word (.docx) files created in docs/rag_sample_word_docs/\n")
